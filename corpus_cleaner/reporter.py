"""Report generator — produces Excel, text, and JSON summaries.

Generates:
- _report.xlsx with 6 worksheets
- _summary.txt human-readable summary
- _original_manifest.json origin → destination mapping
"""

from __future__ import annotations

import json
import re
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# OpenXL / XML illegal characters: U+0000-U+0008, U+000B-U+000C, U+000E-U+001F
_RE_ILLEGAL_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")

HEADER_FONT = Font(bold=True, size=11)
HEADER_FILL = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)

LANGUAGES = ("AR", "RU", "FR", "EN", "ID")
BRACKETS = ("S", "M", "L")


def generate_report(results: dict, output_path: str) -> None:
    """Generate cleaning report as Excel, text summary, and JSON manifest.

    Args:
        results: Pipeline results dict with keys:
            files, contaminated, duplicates, bad_files, stats.
        output_path: Directory where report files will be written.
    """
    out = Path(output_path)
    out.mkdir(parents=True, exist_ok=True)

    _write_excel(results, out / "_report.xlsx")
    _write_summary(results, out / "_summary.txt")
    _write_manifest(results, out / "_original_manifest.json")


def _write_excel(results: dict, path: Path) -> None:
    wb = Workbook()
    wb.remove(wb.active)  # type: ignore[arg-type]

    files: list[dict] = results.get("files", [])
    contaminated: list[dict] = results.get("contaminated", [])
    duplicates: list[dict] = results.get("duplicates", [])
    bad_files: list[dict] = results.get("bad_files", [])

    _sheet_summary(wb, files, contaminated, duplicates, bad_files)
    _sheet_file_details(wb, files)
    _sheet_contaminated(wb, contaminated)
    _sheet_duplicates(wb, duplicates)
    _sheet_bad_files(wb, bad_files)
    _sheet_length_distribution(wb, files)

    wb.save(str(path))


# ── Per-sheet builders ───────────────────────────────────────────────────────


def _sheet_summary(
    wb: Workbook,
    files: list[dict],
    contaminated: list[dict],
    duplicates: list[dict],
    bad_files: list[dict],
) -> None:
    ws = wb.create_sheet("Summary")

    by_lang: dict[str, int] = {}
    for f in files:
        lang = f.get("primary_lang", "")
        by_lang[lang] = by_lang.get(lang, 0) + 1

    by_status: dict[str, int] = {}
    for f in files:
        st = f.get("status", "ok")
        by_status[st] = by_status.get(st, 0) + 1

    by_lang_bracket: dict[tuple[str, str], int] = defaultdict(int)
    for f in files:
        lang = f.get("primary_lang", "")
        bracket = f.get("length_bracket", "")
        by_lang_bracket[(lang, bracket)] += 1

    row = 1
    _write_header_row(ws, row, ["Metric", "Value"])
    row += 1

    rows_data = [
        ("Total files processed", str(len(files))),
        ("Contaminated files", str(len(contaminated))),
        ("Duplicate groups", str(len(duplicates))),
        ("Bad files", str(len(bad_files))),
        ("", ""),
    ]
    for label, val in rows_data:
        ws.cell(row=row, column=1, value=_sanitize_cell(label))
        ws.cell(row=row, column=2, value=_sanitize_cell(val))
        row += 1

    ws.cell(row=row, column=1, value="By Language").font = HEADER_FONT
    row += 1
    # Dynamically determine languages from data (not hardcoded).
    sorted_langs = sorted(set(by_lang.keys()) | set(LANGUAGES))
    for lang in sorted_langs:
        ws.cell(row=row, column=1, value=_sanitize_cell(lang))
        ws.cell(row=row, column=2, value=_sanitize_cell(by_lang.get(lang, 0)))
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="By Status").font = HEADER_FONT
    row += 1
    status_order = ["ok", "table", "contaminated", "duplicate", "empty", "bad", "html_noisy", "image_only"]
    for st in status_order:
        if st in by_status:
            ws.cell(row=row, column=1, value=_sanitize_cell(st))
            ws.cell(row=row, column=2, value=_sanitize_cell(by_status[st]))
            row += 1

    row += 1
    ws.cell(row=row, column=1, value="Language × Bracket").font = HEADER_FONT
    row += 1
    _write_header_row(ws, row, ["Language", "Bracket", "Count"])
    row += 1
    for lang in sorted_langs:
        for bracket in BRACKETS:
            cnt = by_lang_bracket.get((lang, bracket), 0)
            ws.cell(row=row, column=1, value=_sanitize_cell(lang))
            ws.cell(row=row, column=2, value=_sanitize_cell(bracket))
            ws.cell(row=row, column=3, value=_sanitize_cell(cnt))
            row += 1

    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 14


def _sheet_file_details(wb: Workbook, files: list[dict]) -> None:
    ws = wb.create_sheet("FileDetails")

    columns = [
        "source_folder",
        "original_name",
        "new_name",
        "category",
        "primary_lang",
        "detected_lang",
        "line_count",
        "raw_char_count",
        "net_char_count",
        "length_bracket",
        "table_ratio",
        "contamination_score",
        "status",
        "reason",
    ]
    _write_header_row(ws, 1, columns)

    for i, f in enumerate(files):
        row = i + 2
        for j, col in enumerate(columns):
            val = f.get(col, "")
            ws.cell(row=row, column=j + 1, value=_sanitize_cell(val))

    _auto_width(ws, len(columns))


def _sheet_contaminated(wb: Workbook, contaminated: list[dict]) -> None:
    ws = wb.create_sheet("ContaminatedFiles")
    columns = ["filename", "score", "sample_lines"]
    _write_header_row(ws, 1, columns)

    for i, item in enumerate(contaminated):
        row = i + 2
        for j, col in enumerate(columns):
            ws.cell(row=row, column=j + 1, value=_sanitize_cell(item.get(col, "")))

    _auto_width(ws, len(columns))
    ws.column_dimensions[get_column_letter(3)].width = 80


def _sheet_duplicates(wb: Workbook, duplicates: list[dict]) -> None:
    ws = wb.create_sheet("Duplicates")
    columns = ["kept_name", "duplicate_names", "similarity"]
    _write_header_row(ws, 1, columns)

    for i, item in enumerate(duplicates):
        row = i + 2
        ws.cell(row=row, column=1, value=_sanitize_cell(item.get("kept_name", "")))
        dup_names = item.get("duplicate_names", [])
        ws.cell(row=row, column=2, value=_sanitize_cell(", ".join(dup_names) if dup_names else ""))
        ws.cell(row=row, column=3, value=_sanitize_cell(item.get("similarity", "")))

    _auto_width(ws, len(columns))
    ws.column_dimensions[get_column_letter(1)].width = 40
    ws.column_dimensions[get_column_letter(2)].width = 60


def _sheet_bad_files(wb: Workbook, bad_files: list[dict]) -> None:
    ws = wb.create_sheet("BadFiles")
    columns = ["filename", "error"]
    _write_header_row(ws, 1, columns)

    for i, item in enumerate(bad_files):
        row = i + 2
        for j, col in enumerate(columns):
            ws.cell(row=row, column=j + 1, value=_sanitize_cell(item.get(col, "")))

    _auto_width(ws, len(columns))
    ws.column_dimensions[get_column_letter(1)].width = 50
    ws.column_dimensions[get_column_letter(2)].width = 60


def _sheet_length_distribution(wb: Workbook, files: list[dict]) -> None:
    ws = wb.create_sheet("LengthDistribution")

    counts: dict[tuple[str, str], int] = defaultdict(int)
    lang_totals: dict[str, int] = defaultdict(int)
    for f in files:
        lang = f.get("primary_lang", "")
        bracket = f.get("length_bracket", "")
        counts[(lang, bracket)] += 1
        lang_totals[lang] += 1

    cols = ["Language", "S", "M", "L", "S%", "M%", "L%", "Total"]
    _write_header_row(ws, 1, cols)

    sorted_langs = sorted(set(lang_totals.keys()) | set(LANGUAGES))
    for i, lang in enumerate(sorted_langs):
        row = i + 2
        total = lang_totals.get(lang, 0)
        s_cnt = counts.get((lang, "S"), 0)
        m_cnt = counts.get((lang, "M"), 0)
        l_cnt = counts.get((lang, "L"), 0)

        ws.cell(row=row, column=1, value=_sanitize_cell(lang))
        ws.cell(row=row, column=2, value=_sanitize_cell(s_cnt))
        ws.cell(row=row, column=3, value=_sanitize_cell(m_cnt))
        ws.cell(row=row, column=4, value=_sanitize_cell(l_cnt))

        if total > 0:
            ws.cell(row=row, column=5, value=_sanitize_cell(s_cnt / total))
            ws.cell(row=row, column=6, value=_sanitize_cell(m_cnt / total))
            ws.cell(row=row, column=7, value=_sanitize_cell(l_cnt / total))
        else:
            ws.cell(row=row, column=5, value=_sanitize_cell(0))
            ws.cell(row=row, column=6, value=_sanitize_cell(0))
            ws.cell(row=row, column=7, value=_sanitize_cell(0))

        ws.cell(row=row, column=8, value=_sanitize_cell(total))

    for row in range(2, 2 + len(sorted_langs)):
        for col in (5, 6, 7):
            ws.cell(row=row, column=col).number_format = "0.0%"

    _auto_width(ws, len(cols))


# ── Helpers ──────────────────────────────────────────────────────────────────


def _sanitize_cell(value):
    """Strip characters illegal in OpenXML from string values.

    OpenXML forbids U+0000-U+0008, U+000B-U+000C, U+000E-U+001F.
    Numeric values pass through unchanged; None becomes empty string.
    Lists and dicts are converted to their string representation.
    """
    if value is None:
        return ""
    if isinstance(value, (list, dict)):
        return str(value)
    if not isinstance(value, str):
        return value
    return _RE_ILLEGAL_XML.sub("", value)


def _write_header_row(ws, row: int, columns: list[str]) -> None:
    for j, col_name in enumerate(columns):
        cell = ws.cell(row=row, column=j + 1, value=col_name)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGNMENT


def _auto_width(ws, num_cols: int, max_width: int = 50) -> None:
    for col_idx in range(1, num_cols + 1):
        header = str(ws.cell(row=1, column=col_idx).value or "")
        ws.column_dimensions[get_column_letter(col_idx)].width = min(
            len(header) + 6, max_width
        )


def _write_summary(results: dict, path: Path) -> None:
    files: list[dict] = results.get("files", [])
    contaminated: list[dict] = results.get("contaminated", [])
    duplicates: list[dict] = results.get("duplicates", [])
    bad_files: list[dict] = results.get("bad_files", [])

    by_lang: dict[str, int] = {}
    by_status: dict[str, int] = {}
    by_lang_bracket: dict[tuple[str, str], int] = defaultdict(int)
    lang_totals: dict[str, int] = defaultdict(int)

    for f in files:
        lang = f.get("primary_lang", "")
        bracket = f.get("length_bracket", "")
        st = f.get("status", "ok")
        by_lang[lang] = by_lang.get(lang, 0) + 1
        by_status[st] = by_status.get(st, 0) + 1
        by_lang_bracket[(lang, bracket)] += 1
        lang_totals[lang] += 1

    lines: list[str] = []
    lines.append("=" * 60)
    lines.append("  语料清洗报告 (Corpus Cleaning Report)")
    lines.append("=" * 60)
    lines.append("")
    lines.append(f"总处理文件数: {len(files)}")
    lines.append(f"OCR 污染文件: {len(contaminated)}")
    lines.append(f"重复组数: {len(duplicates)}")
    lines.append(f"损坏文件: {len(bad_files)}")
    lines.append("")

    sorted_langs = sorted(set(by_lang.keys()) | set(LANGUAGES))
    lines.append("按语言统计:")
    for lang in sorted_langs:
        lines.append(f"  {lang}: {by_lang.get(lang, 0)}")
    lines.append("")

    lines.append("按状态统计:")
    status_order = ["ok", "table", "contaminated", "duplicate", "empty", "bad", "html_noisy", "image_only"]
    for st in status_order:
        if st in by_status:
            lines.append(f"  {st}: {by_status[st]}")
    lines.append("")

    lines.append("按长度（移动端参考，单位：净字符）:")
    for lang in sorted_langs:
        total = lang_totals.get(lang, 0)
        s_cnt = by_lang_bracket.get((lang, "S"), 0)
        m_cnt = by_lang_bracket.get((lang, "M"), 0)
        l_cnt = by_lang_bracket.get((lang, "L"), 0)
        pct = (l_cnt / total * 100) if total > 0 else 0
        lines.append(
            f"  {lang}: S={s_cnt:<4} M={m_cnt:<4} L={l_cnt:<4}  "
            f"（L 占 {pct:.0f}%，移动端需分页/截断）"
        )
    lines.append("")
    lines.append("=" * 60)

    path.write_text("\n".join(lines), encoding="utf-8")


def _write_manifest(results: dict, path: Path) -> None:
    files: list[dict] = results.get("files", [])
    manifest: dict[str, dict] = {}
    for f in files:
        original = f.get("original_name", "")
        if not original:
            continue
        manifest[original] = {
            "new_name": f.get("new_name", ""),
            "category": f.get("category", ""),
            "primary_lang": f.get("primary_lang", ""),
            "length_bracket": f.get("length_bracket", ""),
            "status": f.get("status", ""),
        }

    path.write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
